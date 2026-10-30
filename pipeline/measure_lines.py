"""Measure how many 18pt body lines actually fit in the figure-page PPT box.

Box geometry mirrors gen_all/kit exactly:
  text box at (PX0=668, y=212) width 552 px (PANEL_W_INNER), font 18pt TNR+SimSun,
  default single leading; hairline reference at y=556 px.
Renders a one-slide PPTX via PowerPoint COM to PNGs, then pixel-scans the
dark-ink bottom row inside the body column and compares against the hairline row.

Re-run this whenever the body font / box / leading changes, and set BODY_MAX
in gen_all.py to the largest N that still prints FITS. Never guess the cap.

Usage:
    python measure_lines.py <paper_workdir> [fig_num]
  <paper_workdir> must contain content.json (the structured figure bodies).
  fig_num defaults to the first figure that has a non-empty slide body.
"""
import json
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from gen_all import _body_lines  # noqa: E402

PX0, BOX_W, BODY_Y, HAIR_Y = 668, 552, 212, 556
EMU = 9525  # 1px = 9525 EMU at 96 dpi


def get_body(wd, fig_num=None):
    """First figure (or the fig_num-th) with a non-empty slide body."""
    wd = Path(wd)
    c = json.loads((wd / "content.json").read_text(encoding="utf-8"))
    figs = c.get("figures", [])
    if fig_num is not None:
        for f in figs:
            if f.get("num") == fig_num:
                for s in f.get("slides", []):
                    if s.get("body"):
                        return s["body"]
        raise SystemExit(f"body not found: fig_num={fig_num}")
    for f in figs:
        for s in f.get("slides", []):
            if s.get("body"):
                return s["body"]
    raise SystemExit("no figure with a non-empty body found in content.json")


def build_body_box(prs, blank, body: str):
    """One slide with the figure-page body box + y556 hairline reference."""
    from pptx.util import Pt, Emu
    from pptx.dml.color import RGBColor
    from lxml import etree

    slide = prs.slides.add_slide(blank)
    line = slide.shapes.add_shape(1, Emu(PX0 * EMU), Emu(HAIR_Y * EMU),
                                  Emu(BOX_W * EMU), Emu(2))
    line.line.color.rgb = RGBColor(0xC9, 0xA2, 0x27)
    line.fill.background()
    line.shadow.inherit = False
    tb = slide.shapes.add_textbox(Emu(PX0 * EMU), Emu(BODY_Y * EMU),
                                  Emu(BOX_W * EMU), Emu((HAIR_Y - BODY_Y + 40) * EMU))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = body
    r.font.size = Pt(18)
    r.font.name = "Times New Roman"
    r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    rPr = r._r.get_or_add_rPr()
    ea = etree.SubElement(rPr, "{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
    ea.set("typeface", "SimSun")
    p.space_before = p.space_after = None
    # match the exported body box: first-line indent 48px (2 chars @18pt)
    pPr = p._p.get_or_add_pPr()
    pPr.set("indent", str(int(48 * 9525)))
    pPr.set("marL", "0")
    return slide


def build_pptx(path: Path, bodies):
    from pptx import Presentation
    from pptx.util import Emu

    prs = Presentation()
    prs.slide_width = Emu(1280 * EMU)
    prs.slide_height = Emu(720 * EMU)
    blank = prs.slide_layouts[6]
    for body in bodies:
        build_body_box(prs, blank, body)
    prs.save(path)


def render_pngs(pptx_path: Path, outdir: Path):
    import win32com.client
    app = win32com.client.Dispatch("PowerPoint.Application")
    try:
        pres = app.Presentations.Open(str(pptx_path), ReadOnly=True, WithWindow=False)
        for i, slide in enumerate(pres.Slides, 1):
            slide.Export(str(outdir / f"slide{i}.png"), "PNG", 1280, 720)
        pres.Close()
    finally:
        app.Quit()


def ink_bottom(png: Path, top: int = 200, bottom: int = 585):
    from PIL import Image
    im = Image.open(png).convert("L")
    w, h = im.size
    scale = w / 1280.0
    px = im.load()
    x0, x1 = int(PX0 * scale), int((PX0 + BOX_W) * scale)
    dark = 100  # #333333 -> 51 gray; gold hairline ~159 and bg 255 stay excluded
    last = None
    for y in range(int(top * scale), min(int(bottom * scale), h)):
        for x in range(x0, x1, 2):
            if px[x, y] < dark:
                last = y
                break
    return (last / scale) if last is not None else None


def main():
    if len(sys.argv) < 2:
        raise SystemExit(__doc__)
    wd = Path(sys.argv[1])
    fig_num = int(sys.argv[2]) if len(sys.argv) > 2 else None
    body = get_body(wd, fig_num)
    est = _body_lines(body)

    out = HERE / "measure_out"
    (out / "png").mkdir(parents=True, exist_ok=True)
    for f in (out / "png").glob("slide*.png"):
        f.unlink()
    pptx = out / "lines_test.pptx"
    build_pptx(pptx, [body])
    render_pngs(pptx, out / "png")

    bottom = ink_bottom(out / "png" / "slide1.png")
    fits = bottom is not None and bottom < HAIR_Y - 1
    print(f"est_lines={est} text_bottom={bottom:.1f}px hairline={HAIR_Y}px "
          f"-> {'FITS' if fits else 'OVERFLOWS'}")
    return fits


if __name__ == "__main__":
    main()
