"""One-command pipeline for one paper workdir (multi-paper batch generation).

Usage:
    <ppt-master-venv-python> run_pipeline.py <workdir> \
        [--name <delivery-basename>] [--desktop-dir <dir>] [--no-export]

The backend scripts root is located in this order:
  1. $PPTMASTER_DIR environment variable (explicit override)
  2. auto-discovery under <home>/AppData/Local/hermes/profiles/*/skills/productivity/ppt-master
     (first one containing scripts/svg_quality_checker.py)

Backend: PPT Master (MIT, https://github.com/hugohe3/ppt-master) is the
recommended toolchain; $PPTMASTER_DIR may point at an equivalent scripts
root from any PPT toolchain that provides the same script names.

Encodes the ordering footguns that used to be tribal knowledge:
  1. project lookup via content.json[paper][ppt_project]; auto-inits (project_manager.py)
     when the dated project dir is gone, and writes the fresh path back into content.json
  2. rm notes/*.md BEFORE total_md_split (its 1:1 check counts stale orphans)
  3. gen_all -> gen_notes -> total_md_split -> compact styles -> compact coords
  4. native-table projects (any kind=="table" figure spec) get
     stamp_native_fallbacks --write AFTER compact and BEFORE the gate
  5. gate: svg_quality_checker --quick-generate --canonical-authoring --stage final --json;
     the report lives at validation/svg_quality_report.json (NOT stdout); abort on blocking>0
  6. export: svg_to_pptx --quick-generate --with-notes --reflow-text [--native-charts-and-tables];
     the fresh export path comes from `ls -t exports/*.pptx`, never from parsing stdout
     (the exporter prints backslash Windows paths that grep mangles)
  7. postproc_justify SRC OUT page_roles.json (restores cx / algn=just / first-line indent)
  8. PPTX verify: slide count == note count, 0 shapes off-slide, <a:tbl> count, media count
  9. optional delivery copy (e.g. Desktop) with busy-file warning

Runs the same steps the skill documents; run it instead of typing the chain.
"""
import argparse
import glob
import json
import os
import pathlib
import re
import shutil
import subprocess
import sys
import zipfile

PY = sys.executable  # must be the ppt-master venv python


def _find_pptmaster() -> pathlib.Path:
    env = os.environ.get("PPTMASTER_DIR")
    if env:
        p = pathlib.Path(env)
        if (p / "scripts" / "svg_quality_checker.py").exists():
            return p
        raise SystemExit(f"PPTMASTER_DIR={env} lacks scripts/svg_quality_checker.py")
    cands = sorted(pathlib.Path.home().glob(
        r"AppData/Local/hermes/profiles/*/skills/productivity/ppt-master"))
    for c in cands:
        if (c / "scripts" / "svg_quality_checker.py").exists():
            return c
    raise SystemExit(
        "ppt-master scripts root not found; set PPTMASTER_DIR to it "
        "(a dir containing scripts/svg_quality_checker.py)")


PPTMASTER = _find_pptmaster()


def run(cmd, **kw):
    print("$", " ".join(str(c) for c in cmd))
    r = subprocess.run([str(c) for c in cmd], **kw)
    if r.returncode != 0:
        raise SystemExit(f"step failed (exit {r.returncode}): {cmd}")


def has_native_tables(content: dict) -> bool:
    return any(f.get("kind") == "table" for f in content.get("figures", []))


def ensure_project(content_path: pathlib.Path, wd: pathlib.Path):
    c = json.loads(content_path.read_text(encoding="utf-8"))
    pdir = pathlib.Path(c["paper"].get("ppt_project", ""))
    if pdir.is_dir() and (pdir / "svg_output").exists():
        return pdir
    import datetime
    # project-name prefix is content-driven (default "figdeck"); a batch can set
    # paper.project_prefix to its own convention without touching this script.
    prefix = c["paper"].get("project_prefix", "figdeck")
    tag = datetime.date.today().strftime("%Y%m%d")
    # init with the dated name the convention uses, then take what was created
    run([PY, PPTMASTER / "scripts" / "project_manager.py", "init",
         f"{prefix}-{wd.name}_ppt169_{tag}", "--format", "ppt169", "--quick-generate"],
        capture_output=True)
    cands = glob.glob(str(PPTMASTER.parent / "projects" / f"{prefix}-{wd.name}_ppt169_*"))
    if not cands:
        raise SystemExit("project_manager init did not create the project dir")
    proj = pathlib.Path(max(cands, key=lambda p: pathlib.Path(p).stat().st_mtime))
    c["paper"]["ppt_project"] = str(proj)
    content_path.write_text(json.dumps(c, ensure_ascii=False, indent=1), encoding="utf-8")
    return proj


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("workdir")
    ap.add_argument("--name", help="delivery file basename (default: <workdir>)")
    ap.add_argument("--desktop-dir", help="copy the delivery PPTX here when done")
    ap.add_argument("--no-export", action="store_true", help="stop after the gate")
    args = ap.parse_args()

    wd = pathlib.Path(args.workdir).resolve()
    content_path = wd / "content.json"
    if not content_path.exists():
        raise SystemExit(f"no content.json in {wd}")
    content = json.loads(content_path.read_text(encoding="utf-8"))
    native = has_native_tables(content)
    out = wd.parent / "output" / f"{args.name or wd.name}_逐图精讲.pptx"

    proj = ensure_project(content_path, wd)
    SHARED = wd.parent  # shared generator scripts live at the batch root (sibling of the workdir)
    if not (SHARED / "gen_all.py").exists():
        raise SystemExit(f"shared scripts (gen_all.py etc.) not found at {SHARED}")
    SK = PPTMASTER / "scripts"

    # 2) notes hygiene
    for f in glob.glob(str(proj / "notes" / "*.md")):
        pathlib.Path(f).unlink()
    # 3) generate
    run([PY, SHARED / "gen_all.py", wd])
    run([PY, SHARED / "gen_notes.py", wd])
    run([PY, SK / "total_md_split.py", proj])
    run([PY, SK / "compact_svg_styles.py", proj / "svg_output", "--inplace"])
    run([PY, SK / "compact_svg_coordinates.py", proj / "svg_output",
         "--inplace", "--keep-native-frames"])
    # 4) native-table fallback binding (AFTER compact rewrites the SVGs)
    if native:
        run([PY, SK / "stamp_native_fallbacks.py", proj / "svg_output", "--write"])
    # 5) gate (the checker exits non-zero on blocking - do NOT abort here; read the report)
    subprocess.run([str(c) for c in [PY, SK / "svg_quality_checker.py", proj,
                  "--quick-generate", "--canonical-authoring", "--stage", "final", "--json"]],
                   capture_output=True)
    rep = json.loads((proj / "validation" / "svg_quality_report.json").read_text(encoding="utf-8"))
    blk = rep["categories"]["blocking"].get("issues", [])
    if blk:
        for i in blk[:10]:
            print("  BLOCKING:", i.get("file"), "|", (i.get("message") or "")[:120])
        raise SystemExit(f"gate blocking={len(blk)} - fix before exporting")
    print(f"gate clean (blocking=0, warnings={rep['summary'].get('warnings')})")
    if args.no_export:
        return

    # 6) export
    exp_flags = ["--quick-generate", "--with-notes", "--reflow-text"]
    if native:
        exp_flags.append("--native-charts-and-tables")
    r = subprocess.run([str(c) for c in [PY, SK / "svg_to_pptx.py", proj, *exp_flags]],
                       capture_output=True, text=True)
    if r.returncode != 0:
        print(r.stdout[-2000:]); print(r.stderr[-2000:]); raise SystemExit("export failed")
    srcs = glob.glob(str(proj / "exports" / "*.pptx"))
    src = max(srcs, key=lambda p: pathlib.Path(p).stat().st_mtime) if srcs else None
    if not src:
        raise SystemExit("no export found in exports/")
    out.parent.mkdir(parents=True, exist_ok=True)
    # 7) postproc (justified + first-line indent + layout-width restore)
    run([PY, SHARED / "postproc_justify.py", src, out, wd / "page_roles.json"])

    # 8) verify
    from pptx import Presentation
    prs = Presentation(str(out))
    W, H = prs.slide_width, prs.slide_height
    notes = bad = 0
    def walk(shapes):
        nonlocal bad
        for s in shapes:
            try: l, t, w, h = s.left, s.top, s.width, s.height
            except Exception: continue
            if l is None: continue
            if l < -1 or t < -1 or l + w > W + 1 or t + h > H + 1:
                bad += 1
            if s.shape_type == 6:
                walk(s.shapes)
    slides = 0
    for s in prs.slides:
        slides += 1
        walk(s.shapes)
        if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip():
            notes += 1
    z = zipfile.ZipFile(str(out))
    tbl = sum(len(re.findall(r"<a:tbl>", z.read(n).decode("utf-8", "ignore")))
              for n in z.namelist() if n.startswith("ppt/slides/"))
    media = len([n for n in z.namelist() if n.startswith("ppt/media/")])
    print(f"VERIFY slides={slides} notes={notes} overflow={bad} native_tbl={tbl} media={media}")
    assert notes == slides, "notes/slide mismatch"
    assert bad == 0, f"{bad} shapes off-slide"
    if native:
        assert tbl > 0, "native tables requested but no <a:tbl> in PPTX"

    # 9) delivery copy
    if args.desktop_dir:
        try:
            shutil.copy2(str(out), args.desktop_dir)
            print(f"copied to {args.desktop_dir}")
        except PermissionError:
            print(f"WARNING: {args.desktop_dir} busy (PPT open?) - file stays in {out.parent}")


if __name__ == "__main__":
    main()
